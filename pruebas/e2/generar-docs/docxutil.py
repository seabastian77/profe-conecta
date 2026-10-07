# -*- coding: utf-8 -*-
"""Ayudas de formato para armar los documentos del Entregable 2 con python-docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_HEX = '0F766E'
TEALCL_HEX = 'D7EDEA'
GRIS_HEX = 'F3F4F6'
NARANJA = RGBColor(0xB4, 0x53, 0x09)
ROJO = RGBColor(0xB9, 0x1C, 0x1C)

SEV_HEX = {'Crítica': '7F1D1D', 'Alta': 'DC2626', 'Media': 'F59E0B', 'Baja': '0EA5E9'}
ESTADO_HEX = {'Aprobado': 'DCFCE7', 'Fallido': 'FEE2E2', 'Bloqueado': 'FEF9C3',
              'Abierto': 'FEE2E2', 'Verificado': 'DCFCE7', 'No ejecutado': 'E5E7EB'}


def _campo(parrafo, instruccion):
    """Inserta un campo de Word (p. ej. PAGE) dentro de un párrafo."""
    run = parrafo.add_run()
    fldBegin = OxmlElement('w:fldChar'); fldBegin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruccion
    fldEnd = OxmlElement('w:fldChar'); fldEnd.set(qn('w:fldCharType'), 'end')
    run._r.append(fldBegin); run._r.append(instr); run._r.append(fldEnd)
    run.font.size = Pt(9); run.font.name = 'Calibri'
    return run


def pie_de_pagina(doc, etiqueta):
    """Pie centrado con la etiqueta del documento y «Página X de Y»."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _A
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.alignment = _A.CENTER
        r = p.add_run(f'{etiqueta}   ·   Página ')
        r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        _campo(p, 'PAGE')
        r2 = p.add_run(' de '); r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        _campo(p, 'NUMPAGES')


def nuevo_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2); s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for nombre, tam in [('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
        st = doc.styles[nombre]
        st.font.name = 'Calibri'; st.font.size = Pt(tam); st.font.bold = True
        st.font.color.rgb = TEAL
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(4)
    return doc


def _set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for lado, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:w'), str(val)); el.set(qn('w:type'), 'dxa')
        m.append(el)
    tcPr.append(m)


def _bordes_tabla(tabla, color='D1D5DB', sz=4):
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def _celda(cell, texto, bold=False, color=None, size=None, align=None, bg=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align: p.alignment = align
    partes = texto.split('\n') if isinstance(texto, str) else [str(texto)]
    for i, parte in enumerate(partes):
        if i: p = cell.add_paragraph()
        if align: p.alignment = align
        r = p.add_run(parte)
        r.font.bold = bold
        if color: r.font.color.rgb = color
        r.font.size = Pt(size if size else 9.5)
    if bg: _set_cell_bg(cell, bg)
    _set_cell_margins(cell)


def tabla(doc, filas, anchos=None, cab=True, fuente=9.5, cab_bg=TEAL_HEX):
    """filas: lista de listas. La primera es encabezado si cab=True."""
    ncol = len(filas[0])
    t = doc.add_table(rows=len(filas), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes_tabla(t)
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            es_cab = cab and i == 0
            _celda(t.cell(i, j), val, bold=es_cab,
                   color=RGBColor(0xFF, 0xFF, 0xFF) if es_cab else None,
                   size=fuente,
                   align=WD_ALIGN_PARAGRAPH.CENTER if es_cab else None,
                   bg=cab_bg if es_cab else (GRIS_HEX if i % 2 == 0 else None))
    if anchos:
        for j, w in enumerate(anchos):
            for i in range(len(filas)):
                t.cell(i, j).width = Cm(w)
    return t


def ficha(doc, pares, ancho_k=4.2, ancho_v=12.5):
    """Tabla de dos columnas clave/valor (para reportes de defecto)."""
    t = doc.add_table(rows=len(pares), cols=2)
    _bordes_tabla(t)
    for i, (k, v) in enumerate(pares):
        _celda(t.cell(i, 0), k, bold=True, size=9.5, bg=TEALCL_HEX)
        _celda(t.cell(i, 1), v, size=9.5)
        t.cell(i, 0).width = Cm(ancho_k); t.cell(i, 1).width = Cm(ancho_v)
    return t


def recuadro(doc, texto, bg=TEALCL_HEX, bold_primero=False):
    t = doc.add_table(rows=1, cols=1)
    _bordes_tabla(t, color=TEAL_HEX, sz=6)
    _celda(t.cell(0, 0), texto, bg=bg, size=10)
    return t


def parrafo(doc, texto, size=11, bold=False, italic=False, space=6, align=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(texto)
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    return p


def codigo(doc, texto):
    """Bloque de código monoespaciado con fondo gris, para los comandos de consola."""
    t = doc.add_table(rows=1, cols=1)
    _bordes_tabla(t, color='D1D5DB', sz=4)
    cell = t.cell(0, 0)
    cell.text = ''
    _set_cell_bg(cell, '1E1E1F')
    _set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    for i, linea in enumerate(texto.split('\n')):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(linea or ' ')
        r.font.name = 'Consolas'; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xE8, 0xEA, 0xED)
    return t


def vineta(doc, texto, nivel=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.8 + nivel * 0.6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(texto); r.font.size = Pt(10.5)
    return p


def numerada(doc, texto):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(texto); r.font.size = Pt(10.5)
    return p


def imagen(doc, ruta, ancho_cm=16):
    doc.add_picture(ruta, width=Cm(ancho_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def salto(doc):
    doc.add_page_break()


def titulo_portada(doc, lineas):
    for texto, size, color, bold in lineas:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(texto); r.font.size = Pt(size); r.font.bold = bold
        if color: r.font.color.rgb = color
