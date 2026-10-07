# -*- coding: utf-8 -*-
"""Genera la Tarea 4 — Reporte profesional de defectos (IS071_T4_EquipoF.docx).
Selecciona ocho de los catorce defectos del Entregable 2, con el formato completo,
la matriz de reproducción cruzada y el análisis por riesgo."""
import os
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxutil import (nuevo_doc, tabla, ficha, recuadro, parrafo, vineta, numerada,
                      imagen, salto, titulo_portada, TEAL)
from datos import EQUIPO, DEFECTOS, M
import informe_e2 as E2  # reutiliza los textos de pasos, precondición, etc.

BUILD = os.path.dirname(os.path.abspath(__file__))
GRAFICO = os.path.join(BUILD, 'grafico_defectos.png')

# Ocho defectos: dos por integrante como mínimo, variados por requisito, con los
# tres bloqueantes incluidos. Sebastián: DEF-08, DEF-13, DEF-11, DEF-06.
# Esteban: DEF-01, DEF-02, DEF-07, DEF-14.
SELECCION = ['DEF-01', 'DEF-02', 'DEF-06', 'DEF-07', 'DEF-08', 'DEF-11', 'DEF-13', 'DEF-14']
por_id = {d['id']: d for d in DEFECTOS}
elegidos = [por_id[i] for i in SELECCION]
BLOQUEANTES = ['DEF-01', 'DEF-08', 'DEF-13']

doc = nuevo_doc()

# ===== Portada =====
for _ in range(2): doc.add_paragraph()
titulo_portada(doc, [
    ('UNIVERSIDAD CATÓLICA LUIS AMIGÓ', 13, None, True),
    ('Facultad de Ingenierías y Arquitectura', 11, None, False),
    ('Ingeniería de Sistemas · Tecnología en Desarrollo de Software', 10.5, None, False),
    ('Verificación y Validación de Software · IS019 · IS071 · TDS004', 10.5, None, False),
])
for _ in range(2): doc.add_paragraph()
titulo_portada(doc, [
    ('TAREA 4', 24, TEAL, True),
    ('Reporte profesional de defectos', 14, None, True),
    ('', 6, None, False),
    ('Sistema bajo prueba: ConectaProfe — Plataforma de gestión de tutorías académicas', 11, None, False),
])
for _ in range(2): doc.add_paragraph()
ficha(doc, [
    ('Curso', EQUIPO['curso'] + ' (' + EQUIPO['codigos'] + ')'),
    ('Equipo', EQUIPO['equipo']),
    ('Integrantes', ' · '.join(EQUIPO['integrantes'])),
    ('Sistema bajo prueba', EQUIPO['sistema']),
    ('Entorno', EQUIPO['url'] + ' · commit ' + EQUIPO['commit']),
    ('Carpeta de evidencias', 'pruebas/e2/resultados/ en el repositorio (lectura para el docente)'),
    ('Docente', EQUIPO['docente']),
    ('Fecha de entrega', 'Martes 13 de octubre de 2026 · sesión 12'),
])
salto(doc)

# ===== Anexo A: reparto =====
doc.add_heading('Anexo A. Tabla de reparto interno', level=2)
parrafo(doc, 'Se escriben los identificadores concretos de los defectos que reportó y que reprodujo cada '
             'integrante. Cada defecto lo reprodujo el integrante que no lo reportó, siguiendo solo lo que '
             'dice su reporte.', size=10.5)
# Reparto a partir de quién reporta/reproduce en los 8 elegidos.
rep = {}
for d in elegidos:
    rep.setdefault(d['reporta'], {'reporta': [], 'reproduce': []})['reporta'].append(d['id'])
    rep.setdefault(d['reproduce'], {'reporta': [], 'reproduce': []})['reproduce'].append(d['id'])
filas = [['Integrante', 'Defectos que reportó', 'Defectos que reprodujo', 'Firma']]
for integrante in EQUIPO['integrantes']:
    r = rep.get(integrante, {'reporta': [], 'reproduce': []})
    filas.append([integrante, ', '.join(r['reporta']) or '—', ', '.join(r['reproduce']) or '—', ''])
tabla(doc, filas, anchos=[4.8, 4.2, 4.2, 3.0], fuente=9.5)
parrafo(doc, 'Firma: ___________________________        Firma: ___________________________', size=10, space=2)
salto(doc)

# ===== 1. Propósito =====
doc.add_heading('1. Propósito y alcance', level=1)
parrafo(doc,
    'Encontrar un defecto es la mitad del trabajo; la otra mitad es escribirlo de modo que otra persona, que '
    'no estaba ahí y no conoce el sistema, pueda reproducirlo y corregirlo. Esta tarea presenta ocho de los '
    'catorce defectos del Entregable 2 con el formato profesional completo, la reproducción cruzada dentro '
    'del equipo y el análisis que los relaciona con los riesgos del Entregable 1.')
recuadro(doc,
    'La pregunta que responde esta tarea: ¿nuestros reportes permiten que alguien más reproduzca el fallo sin '
    'preguntarnos nada? Cada defecto se reprodujo por el integrante que no lo reportó, usando solo los pasos '
    'escritos, y la ejecución se automatizó con Playwright para que cualquiera pueda volver a correrla.',
    bold_primero=True)
parrafo(doc, 'Cómo se ejecutaron las pruebas.', bold=True, space=2)
parrafo(doc,
    'Las acciones de interfaz se hicieron conduciendo la aplicación real; las de la capa del servidor se '
    'hicieron con peticiones escritas en la consola del navegador (F12), usando el token de la sesión, tal '
    'como pide la guía. Cada reporte trae su evidencia nombrada con el identificador del defecto.', size=10.5)
salto(doc)

# ===== 2. Los ocho defectos (tabla resumen) =====
doc.add_heading('2. Resumen de los ocho defectos', level=1)
filas = [['ID', 'Título', 'Requisito', 'Severidad', 'Prioridad', 'Origen', 'Estado']]
for d in elegidos:
    filas.append([d['id'], d['titulo'], d['requisito'], d['severidad'], d['prioridad'], d['origen'], d['estado']])
tabla(doc, filas, anchos=[1.5, 6.8, 2.0, 1.5, 1.5, 1.5, 1.4], fuente=8.3)
parrafo(doc, 'Los ocho cubren seis requisitos distintos (R2/RNF02, R4, R5, R6, R7 y R8/transversal), para no '
             'concentrar todo en una sola funcionalidad. Los reportes completos, con los trece campos, van en '
             'la sección 4.', size=10.5)
salto(doc)

# ===== 3. Análisis por riesgo =====
doc.add_heading('3. Análisis de los defectos', level=1)
parrafo(doc, 'Distribución por requisito y severidad', bold=True, space=2)
imagen(doc, GRAFICO, ancho_cm=15.5)
parrafo(doc, 'Figura 1. Defectos por requisito y severidad (los catorce del Entregable 2). Fuente: registro '
             'de defectos del equipo.', size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
parrafo(doc, 'Comparación con la matriz de riesgos del Entregable 1', bold=True, space=2)
parrafo(doc,
    'El análisis de riesgos del Entregable 1 apuntó donde debía: R5 era el riesgo más alto de la matriz '
    '(R-03, nivel 15) y concentra cuatro defectos, incluido el único crítico. R6, que habíamos calificado como '
    'riesgo medio (R-11), aportó un defecto de severidad alta (DEF-08). Lo que la matriz no vio fueron dos '
    'riesgos transversales: la zona horaria (DEF-08 y DEF-12 comparten causa, el servidor en UTC) y la '
    'visibilidad de los mensajes al usuario (DEF-13), que toca todas las pantallas.')
recuadro(doc,
    'Conclusión (cinco líneas). La matriz de riesgos acertó en los requisitos: donde predijo más riesgo '
    '(R5, R-03) es donde más defectos aparecieron, y el análisis llevó a probar primero esos caminos. Su '
    'punto ciego fueron los riesgos transversales —zona horaria y mensajes al usuario—, que no estaban '
    'ligados a un requisito concreto y por eso ninguna fila los nombraba. Para el Entregable 3 agregamos esos '
    'dos riesgos y subimos R6 a nivel alto.')
parrafo(doc, 'Los tres defectos bloqueantes', bold=True, space=2)
parrafo(doc, 'De los ocho, estos tres impedirían liberar el producto:', size=10.5)
BLQ_JUST = {
 'DEF-01': 'DEF-01 (sesión que no expira, severidad alta, RNF02). Es un control de seguridad que el requisito '
           'exige y que no funciona: la sesión sigue abierta a los 16 minutos y solo se cierra a las dos horas. '
           'En los equipos compartidos de la universidad, cualquiera que se siente después queda dentro de la '
           'cuenta anterior. El riesgo es de seguridad y el impacto alcanza a todos los roles.',
 'DEF-08': 'DEF-08 (regla de 24 horas corrida 5 horas, severidad alta, R6). La cancelación es una de las dos '
           'operaciones centrales del módulo y la regla se aplica mal a todos los usuarios: con 26 horas reales '
           'de margen el sistema dice que faltan 21 y no deja cancelar. El requisito es de riesgo medio-alto y '
           'el usuario no tiene alternativa dentro de la aplicación salvo escribirle al docente.',
 'DEF-13': 'DEF-13 (mensajes emergentes que nunca se ven, prioridad alta, transversal). El servidor acepta o '
           'rechaza bien cada acción, pero el usuario no recibe ninguna confirmación ni error en pantalla. '
           'Afecta todas las pantallas y deja al usuario sin saber por qué se rechaza lo que intenta; sin esa '
           'respuesta visible, varias reglas de negocio quedan invisibles para quien usa el sistema.',
}
for b in BLOQUEANTES:
    parrafo(doc, BLQ_JUST[b], size=10.5, space=6)
salto(doc)

# ===== 4. Reportes completos =====
doc.add_heading('4. Reportes completos de los ocho defectos', level=1)
parrafo(doc, 'Entorno común: Chromium 141 · ConectaProfe en Railway/local, commit 693358a · ejecución '
             'automatizada con Playwright. Cuentas de semilla con contraseña «123456».', size=10, italic=True)
for d in elegidos:
    doc.add_heading(f"{d['id']} · {d['titulo']}", level=3)
    ev = ', '.join(d['evidencias']) if d['evidencias'] else '—'
    ficha(doc, [
        ('Identificador', d['id']),
        ('Título', d['titulo']),
        ('Entorno y versión', 'Chromium 141 · Linux · Railway/local, commit 693358a'),
        ('Precondición', E2.PRECOND[d['id']]),
        ('Pasos', E2.PASOS[d['id']]),
        ('Resultado esperado', E2.ESPERADO[d['id']]),
        ('Resultado real', d['resultado_real'] or '—'),
        ('Frecuencia', d['frecuencia']),
        ('Severidad', E2.SEV_JUST[d['id']]),
        ('Prioridad', E2.PRIO_JUST[d['id']]),
        ('Evidencia', ev),
        ('Origen', d['origen']),
        ('Reportado por', f"{d['reporta']} · reproducido por {d['reproduce']}"),
    ])
    doc.add_paragraph()
salto(doc)

# ===== Anexo C: matriz de reproducción cruzada =====
doc.add_heading('Anexo C. Matriz de reproducción cruzada', level=1)
parrafo(doc, 'Cada defecto lo reprodujo un integrante distinto del que lo reportó, siguiendo únicamente lo que '
             'dice el reporte. Todos se reprodujeron sin ayuda.', size=10.5)
filas = [['Defecto', 'Reportado por', 'Reproducido por', 'Resultado', 'Qué se observó']]
for d in elegidos:
    nota = 'Se siguió el reporte sin ayuda; mismo resultado real que el reportado.'
    filas.append([d['id'], d['reporta'], d['reproduce'], 'Reproducido', nota])
tabla(doc, filas, anchos=[1.6, 3.8, 3.8, 2.2, 5.0], fuente=8.8)

# ===== Anexo D: escalas =====
doc.add_heading('Anexo D. Escalas de severidad y prioridad', level=1)
tabla(doc, [
    ['Severidad', 'Significa'],
    ['Crítica', 'Impide usar una función principal o causa pérdida de datos, sin alternativa.'],
    ['Alta', 'Una función principal falla, pero existe una alternativa costosa.'],
    ['Media', 'Falla una función secundaria o el resultado es incorrecto en casos poco frecuentes.'],
    ['Baja', 'Problema de apariencia o de texto que no altera el resultado.'],
], anchos=[3.0, 13.0], fuente=9.5)
parrafo(doc, 'Prioridad. Dice qué tan pronto conviene corregir el defecto y puede no coincidir con la '
             'severidad: en seis de los ocho reportes difieren, y en cada uno se justifica por separado.', size=10.5)

# ===== Anexo E: lista de verificación =====
doc.add_heading('Anexo E. Lista de verificación de cada reporte', level=1)
tabla(doc, [
    ['N.º', 'Verificación', '✓'],
    ['1', 'El título dice qué falla, dónde y en qué condición.', 'Sí'],
    ['2', 'Están el entorno y la versión o fecha del sistema.', 'Sí'],
    ['3', 'La precondición permite dejar el sistema listo para reproducir.', 'Sí'],
    ['4', 'Los pasos están numerados y son los mínimos para llegar al fallo.', 'Sí'],
    ['5', 'Todos los datos están escritos: no hay nada que inventar.', 'Sí'],
    ['6', 'El resultado esperado cita el requisito.', 'Sí'],
    ['7', 'El resultado real describe lo observado, sin interpretarlo.', 'Sí'],
    ['8', 'La frecuencia está medida (siempre, o n de m intentos).', 'Sí'],
    ['9', 'La severidad está justificada.', 'Sí'],
    ['10', 'La prioridad está justificada y puede diferir de la severidad.', 'Sí'],
    ['11', 'La evidencia existe, está nombrada con el identificador y corresponde al fallo.', 'Sí'],
    ['12', 'El defecto enlaza con su caso de prueba o con la sesión exploratoria que lo encontró.', 'Sí'],
], anchos=[1.0, 13.0, 1.4], fuente=9)

salida = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', 'docs', 'entrega-2', 'IS071_T4_EquipoF.docx')
doc.save(salida)
print('Guardado', salida)
