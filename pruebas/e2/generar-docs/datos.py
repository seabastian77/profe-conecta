# -*- coding: utf-8 -*-
"""Datos del Entregable 2 y la Tarea 4.

La fuente es lo que el equipo ejecuta a mano en Railway:
  - docs/entrega-2/E2_Registro_Ejecucion_EquipoF.xlsx (columnas amarillas: fecha, navegador,
    estado, resultado real, intentos, evidencia, reproducción cruzada y sesiones exploratorias);
  - las capturas en pruebas/e2/evidencias-equipo/.

De pruebas/e2/resultados/ (la corrida del script sobre una copia local) solo se toman las
definiciones de los casos (título, requisito, prioridad, resultado esperado); sus capturas y
resultados son referencia y no entran a los documentos.
"""
import datetime
import json
import os
import re

import openpyxl

from textos import DEFECTOS_BASE, EVIDENCIA_SUGERIDA, MISIONES

BUILD = os.path.dirname(os.path.abspath(__file__))
RAIZ_E2 = os.path.normpath(os.path.join(BUILD, '..'))
REGISTRO = os.path.normpath(os.path.join(BUILD, '..', '..', '..', 'docs', 'entrega-2',
                                         'E2_Registro_Ejecucion_EquipoF.xlsx'))
CAPTURAS = os.path.join(RAIZ_E2, 'evidencias-equipo')

EQUIPO = {
    'curso': 'Verificación y Validación de Software',
    'codigos': 'IS019 · IS071 · TDS004',
    'docente': 'Camilo Rendón Vinasco',
    'periodo': '2026-2 · Sede Medellín',
    'equipo': 'F',
    'integrantes': ['Sebastián González González', 'Esteban Palencia'],
    'sistema': 'ConectaProfe · plataforma de gestión de tutorías académicas',
    'url': 'https://profe-conecta-production-e40c.up.railway.app',
    'commit': '693358a (rama main)',
    'entrega': 'Martes 13 de octubre de 2026 · sesión 12',
}
PENDIENTE = '____'

# --- Definición de los casos (del diseño de la suite) --------------------------
with open(os.path.join(RAIZ_E2, 'resultados', 'ciclo-1', 'resultados.json'), encoding='utf-8') as _f:
    CASOS = {c['id']: {k: c[k] for k in ('id', 'titulo', 'requisito', 'riesgo', 'tecnica', 'prioridad',
                                          'responsable', 'esperado')}
             for c in json.load(_f)['casos']}


def _texto(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%d/%m/%Y')
    return str(v).strip()


def _lista(v):
    """Nombres separados por coma, punto y coma, salto de línea o « y »; sin repetidos."""
    vistos = []
    for x in re.split(r'[,;\n]|\s+y\s+', _texto(v)):
        x = x.strip()
        if x and x not in vistos:
            vistos.append(x)
    return vistos


AVISOS = []   # lo que no se pudo entender del registro; sale en el aviso amarillo


def _estado(valor, opciones, donde):
    """Normaliza lo que escribió el equipo («aprobado», «Fallo»…); avisa si no se reconoce."""
    t = _texto(valor)
    if not t:
        return ''
    base = t.strip().lower()
    for o in opciones:
        if base == o.lower() or (len(base) >= 4 and o.lower().startswith(base)):
            return o
    if base in ('fallo', 'falló', 'falla'):
        return 'Fallido'
    AVISOS.append(f'{donde}: «{t}» no es un estado válido ({", ".join(opciones)})')
    return ''


ESTADOS_EJEC = ('Aprobado', 'Fallido', 'Bloqueado')
RESULTADOS_REPRO = ('Reproducido', 'Con ayuda', 'No reproducido')


# --- Capturas del equipo ----------------------------------------------------------
_ARCHIVOS = {}
if os.path.isdir(CAPTURAS):
    for _a in os.listdir(CAPTURAS):
        if _a.lower().endswith(('.png', '.jpg', '.jpeg')):
            _ARCHIVOS[os.path.splitext(_a)[0].lower()] = os.path.join(CAPTURAS, _a)


def ruta_captura(nombre):
    """Ruta de una captura del equipo; acepta .png, .jpg o .jpeg con el mismo nombre."""
    return _ARCHIVOS.get(os.path.splitext(nombre)[0].lower())


# --- Registro del equipo ----------------------------------------------------------
def _hoja(wb, nombre):
    if nombre not in wb.sheetnames:
        return []
    ws = wb[nombre]
    cab = [_texto(c.value) for c in ws[1]]
    filas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in fila):
            continue
        filas.append({cab[j]: fila[j] for j in range(min(len(cab), len(fila)))})
    return filas


if not os.path.exists(REGISTRO):
    raise SystemExit(f'Falta el registro {REGISTRO}; créelo con plantilla_registro.py')
_wb = openpyxl.load_workbook(REGISTRO, data_only=True)

EJEC = []
for f in _hoja(_wb, 'Registro de ejecución'):
    caso = _texto(f.get('Caso'))
    if not caso.startswith('CP-'):
        continue
    evid = _lista(f.get('Evidencia'))
    e = dict(
        caso=caso, ciclo=int(f.get('Ciclo') or 1), que=_texto(f.get('Qué prueba')), ejecuta=_texto(f.get('Lo ejecuta')),
        navegador=_texto(f.get('Navegador y versión')), so=_texto(f.get('Sistema operativo')),
        fecha=_texto(f.get('Fecha')), real=_texto(f.get('Resultado real')),
        evidencias=evid, defecto=_texto(f.get('Defecto')).upper().replace(' ', ''),
    )
    e['estado'] = _estado(f.get('Estado'), ESTADOS_EJEC, f"{caso} ciclo {e['ciclo']}")
    e['rutas'] = [ruta_captura(x) for x in evid]
    e['hecho'] = bool(e['estado'])
    EJEC.append(e)

# Autor y fecha de cada captura de ejecución, para los pies de figura de los defectos.
AUTOR_CAPTURA = {n: (e['ejecuta'], e['fecha']) for e in EJEC for n in e['evidencias']}

_def_eq = {_texto(f.get('ID')): f for f in _hoja(_wb, 'Defectos')}
DEFECTOS = []
for base in DEFECTOS_BASE:
    f = _def_eq.get(base['id'], {})
    d = dict(base)
    d.update(fecha=_texto(f.get('Fecha')), navegador=_texto(f.get('Navegador y versión')),
             so=_texto(f.get('Sistema operativo')), intentos=_texto(f.get('Intentos')),
             resultado_real=_texto(f.get('Resultado real')),
             evidencias=_lista(f.get('Evidencia')) or list(EVIDENCIA_SUGERIDA[base['id']]))
    d['rutas'] = [ruta_captura(x) for x in d['evidencias']]
    d['hecho'] = bool(d['fecha'] and d['resultado_real'])
    d['estado_equipo'] = _estado(f.get('Estado'), ('Abierto', 'Verificado', 'No reproducido'), base['id'])
    DEFECTOS.append(d)

REPRO = {}
for f in _hoja(_wb, 'Reproducción cruzada'):
    did = _texto(f.get('Defecto'))
    REPRO[did] = dict(fecha=_texto(f.get('Fecha')), falto=_texto(f.get('Qué faltó o qué se observó')),
                      resultado=_estado(f.get('Resultado'), RESULTADOS_REPRO, f'Reproducción de {did}'))

# Estado de cada defecto según lo que el equipo ya hizo, nunca por adelantado:
# - DEF-05 está corregido en el código; pasa a «Verificado» cuando quien lo reporta repite sus pasos en Railway.
# - Los demás quedan «Por reproducir» hasta que el otro integrante lo intente; «Abierto» si le salió
#   (solo o con ayuda) y «No reproducido» si no le salió.
for d in DEFECTOS:
    r = REPRO.get(d['id'], {}).get('resultado')
    if d['id'] == 'DEF-05':
        d['estado'] = 'Verificado' if d['hecho'] and d['estado_equipo'] == 'Verificado' else 'Corregido, por verificar'
    elif d['estado_equipo'] == 'No reproducido' or r == 'No reproducido':
        d['estado'] = 'No reproducido'
    elif r in ('Reproducido', 'Con ayuda'):
        d['estado'] = 'Abierto'
    else:
        d['estado'] = 'Por reproducir'

for e in EJEC:
    if e['estado'] == 'Fallido' and not e['defecto']:
        AVISOS.append(f"{e['caso']} ciclo {e['ciclo']}: está Fallido y le falta el defecto")
    for n, r in zip(e['evidencias'], e['rutas']):
        if r and os.path.splitext(n)[1].lower() != os.path.splitext(r)[1].lower():
            e['evidencias'][e['evidencias'].index(n)] = os.path.basename(r)

SESIONES = {}
NOTAS = {'SE-01': [], 'SE-02': []}
for f in _hoja(_wb, 'Sesiones exploratorias'):
    sid = _texto(f.get('Sesión'))
    if sid:
        SESIONES[sid] = dict(integrante=_texto(f.get('Integrante')), mision=_texto(f.get('Misión')) or MISIONES.get(sid, ''),
                             fecha=_texto(f.get('Fecha')), duracion=_texto(f.get('Duración (min)')),
                             defectos=_texto(f.get('Defectos encontrados')), preguntas=_texto(f.get('Preguntas y riesgos')),
                             distribucion=_texto(f.get('Distribución del tiempo')))
for f in _hoja(_wb, 'Notas de sesión'):
    sid = _texto(f.get('Sesión'))
    minuto, hice, vi = _texto(f.get('Minuto')), _texto(f.get('Qué hice')), _texto(f.get('Qué vi'))
    if sid in NOTAS and (minuto or hice or vi):
        NOTAS[sid].append((minuto, hice, vi))


# --- Métricas (solo con lo que el equipo ya ejecutó) -----------------------------
def _contar(ciclo, estado):
    return sum(1 for e in EJEC if e['ciclo'] == ciclo and e['estado'] == estado)


def metricas():
    m = {}
    for c in (1, 2):
        m[f'aprob{c}'] = _contar(c, 'Aprobado')
        m[f'fall{c}'] = _contar(c, 'Fallido')
        m[f'bloq{c}'] = _contar(c, 'Bloqueado')
        m[f'ejec{c}'] = m[f'aprob{c}'] + m[f'fall{c}']
        m[f'dis{c}'] = sum(1 for e in EJEC if e['ciclo'] == c)
        m[f'pend{c}'] = sum(1 for e in EJEC if e['ciclo'] == c and not e['hecho'])
    alta = [e for e in EJEC if e['ciclo'] == 1 and CASOS.get(e['caso'], {}).get('prioridad') == 'Alta']
    m['alta'] = len(alta)
    m['alta_ejec'] = sum(1 for e in alta if e['estado'] in ('Aprobado', 'Fallido'))
    m['total_def'] = len(DEFECTOS)
    m['def_hechos'] = sum(1 for d in DEFECTOS if d['hecho'] and d['navegador'] and d['so'] and d['intentos'])
    m['repro_hechas'] = sum(1 for d in DEFECTOS if REPRO.get(d['id'], {}).get('resultado'))
    m['reproducidos'] = sum(1 for d in DEFECTOS if REPRO.get(d['id'], {}).get('resultado') in ('Reproducido', 'Con ayuda'))
    m['capturas'] = sum(1 for e in EJEC for r in e['rutas'] if r)
    m['capturas_esperadas'] = sum(len(e['evidencias']) for e in EJEC)
    m['capturas_def_faltan'] = sum(1 for d in DEFECTOS for n, r in zip(d['evidencias'], d['rutas'])
                                   if not r and n not in AUTOR_CAPTURA)
    m['ejec_incompletas'] = sum(1 for e in EJEC if e['hecho'] and not (e['fecha'] and e['navegador'] and e['so']))
    m['sesiones_faltan'] = sum(1 for s in ('SE-01', 'SE-02')
                               if not (SESIONES.get(s, {}).get('fecha') and SESIONES.get(s, {}).get('duracion') and NOTAS.get(s)))
    return m


M = metricas()
COMPLETO = (M['pend1'] == 0 and M['pend2'] == 0 and M['ejec_incompletas'] == 0 and M['def_hechos'] == len(DEFECTOS)
            and M['repro_hechas'] == len(DEFECTOS) and M['capturas'] == M['capturas_esperadas']
            and M['capturas_def_faltan'] == 0 and M['sesiones_faltan'] == 0 and not AVISOS)


def v(x, raya=PENDIENTE):
    """El dato del equipo o una raya para llenar."""
    return x if x not in (None, '') else raya


def entorno(navegador, so):
    """Entorno de una ejecución manual en Railway."""
    partes = [p for p in (navegador, so) if p]
    return 'Railway (commit 693358a) · ' + (' · '.join(partes) if partes else
                                            'navegador y versión: ____ · sistema operativo: ____')


if __name__ == '__main__':
    for a in AVISOS:
        print('AVISO:', a)
    print('Ejecuciones:', len(EJEC), '· pendientes', M['pend1'] + M['pend2'])
    print('Defectos con datos del equipo:', M['def_hechos'], 'de', len(DEFECTOS))
    print('Reproducciones cruzadas:', M['repro_hechas'], '· capturas', M['capturas'], 'de', M['capturas_esperadas'])
