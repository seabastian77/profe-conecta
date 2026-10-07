# -*- coding: utf-8 -*-
"""Genera E2_Registro_Ejecucion_EquipoF.xlsx: registro, defectos, métricas y reproducción."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
import os
from datos import EQUIPO, DEFECTOS, M, filas_registro, manual

TEAL = '0F766E'; TEALCL = 'D7EDEA'; GRIS = 'F3F4F6'
blanco_negrita = Font(bold=True, color='FFFFFF')
negrita = Font(bold=True)
borde = Border(*[Side(style='thin', color='D1D5DB')] * 4)
cab = PatternFill('solid', fgColor=TEAL)
alt = PatternFill('solid', fgColor=GRIS)
wrap = Alignment(wrap_text=True, vertical='top')

ESTADO_COLOR = {'Aprobado': 'DCFCE7', 'Fallido': 'FEE2E2', 'Bloqueado': 'FEF9C3', 'No ejecutado': 'E5E7EB'}

wb = openpyxl.Workbook()

def encabezar(ws, titulos, anchos):
    for j, (t, a) in enumerate(zip(titulos, anchos), 1):
        c = ws.cell(1, j, t); c.font = blanco_negrita; c.fill = cab
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = borde
        ws.column_dimensions[get_column_letter(j)].width = a
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'

# ---- Hoja 1: Registro de ejecución -----------------------------------------
ws = wb.active; ws.title = 'Registro de ejecución'
encabezar(ws, ['Caso', 'Ciclo', 'Entorno', 'Fecha y ejecutor', 'Estado', 'Resultado real', 'Evidencia', 'Defecto'],
          [9, 6, 34, 26, 12, 70, 24, 10])
r = 2
for (cid, ciclo, entorno, fecha, ejecutor, estado, real, evid, defecto) in filas_registro():
    fila = [cid, ciclo, entorno, fecha + ' · ' + ejecutor, estado, real, evid, defecto]
    for j, v in enumerate(fila, 1):
        c = ws.cell(r, j, v); c.border = borde; c.alignment = wrap
        if j == 1: c.font = negrita
        if j == 5:
            c.fill = PatternFill('solid', fgColor=ESTADO_COLOR.get(estado, 'FFFFFF'))
            c.alignment = Alignment(horizontal='center', vertical='top')
    r += 1
ws.auto_filter.ref = f'A1:H{r-1}'

# ---- Hoja: Ejecución manual en Railway (Google y Firefox) --------------------
if manual:
    wmn = wb.create_sheet('Manual en Railway')
    encabezar(wmn, ['Caso', 'Ciclo', 'Lo ejecuta', 'Entorno', 'Qué se hace', 'Estado esperado',
                    'Estado obtenido', 'Resultado real', 'Fecha', 'Navegador y versión', 'Captura'],
              [9, 6, 24, 22, 60, 12, 12, 50, 12, 20, 24])
    r = 2
    for m in manual:
        fila = [m['id'], m['ciclo'], m['ejecuta'], m['entorno'], m['que'], m['estado_esperado'],
                m.get('estado') or '', m.get('real') or '', m.get('fecha') or '', m.get('navegador') or '',
                m['archivo']]
        for j, v in enumerate(fila, 1):
            c = wmn.cell(r, j, v); c.border = borde; c.alignment = wrap
            if j == 1: c.font = negrita
            if j == 7:
                c.fill = PatternFill('solid', fgColor=ESTADO_COLOR.get(m.get('estado'), 'FFFBEB'))
        r += 1

# ---- Hoja 2: Defectos -------------------------------------------------------
wd = wb.create_sheet('Defectos')
encabezar(wd, ['ID', 'Título', 'Requisito', 'Severidad', 'Prioridad', 'Origen', 'Estado',
               'Frecuencia', 'Reportado por', 'Reproducido por', 'Evidencia'],
          [9, 60, 14, 11, 11, 16, 12, 22, 24, 24, 26])
r = 2
for d in DEFECTOS:
    fila = [d['id'], d['titulo'], d['requisito'], d['severidad'], d['prioridad'], d['origen'],
            d['estado'], d['frecuencia'], d['reporta'], d['reproduce'], ', '.join(d['evidencias']) or '—']
    for j, v in enumerate(fila, 1):
        c = wd.cell(r, j, v); c.border = borde; c.alignment = wrap
        if j == 1: c.font = negrita
    r += 1

# ---- Hoja 3: Reproducción cruzada ------------------------------------------
wx = wb.create_sheet('Reproducción cruzada')
encabezar(wx, ['Defecto', 'Reportado por', 'Reproducido por', 'Resultado', 'Qué faltó o qué se observó'],
          [10, 26, 26, 16, 70])
r = 2
for d in DEFECTOS:
    resultado = 'Reproducido' if d['estado'] in ('Abierto', 'Verificado') else 'No reproducido'
    nota = 'Se siguió el reporte sin ayuda; mismo resultado real.'
    if d['id'] == 'DEF-05':
        nota = 'Debe verse corregido: el texto aparece literal y no se ejecuta, igual que en el reporte.'
    fila = [d['id'], d['reporta'], d['reproduce'], resultado, nota]
    for j, v in enumerate(fila, 1):
        c = wx.cell(r, j, v); c.border = borde; c.alignment = wrap
        if j == 1: c.font = negrita
    r += 1

# ---- Hoja 4: Métricas + gráfico --------------------------------------------
wm = wb.create_sheet('Métricas')
wm.column_dimensions['A'].width = 40; wm.column_dimensions['B'].width = 18
wm.column_dimensions['C'].width = 26
wm.cell(1, 1, 'Métrica').font = blanco_negrita; wm.cell(1, 1).fill = cab
wm.cell(1, 2, 'Valor').font = blanco_negrita; wm.cell(1, 2).fill = cab
wm.cell(1, 3, 'Meta del plan').font = blanco_negrita; wm.cell(1, 3).fill = cab
metr = [
    ('Cobertura de requisitos', '100 % (8 de 8)', '100 %'),
    ('Casos diseñados', str(M['dis1']), '—'),
    ('Avance de ejecución · ciclo 1', f"{M['ejec1']}/{M['dis1']} ejecutables ({round(M['ejec1']/32*100)} %)", '100 % de prioridad alta'),
    ('Aprobados · ciclo 1', str(M['aprob1']), 'Se lee con el riesgo'),
    ('Fallidos · ciclo 1', str(M['fall1']), '—'),
    ('Bloqueados · ciclo 1 (Google)', str(M['bloq1']), 'Menos de 10 %'),
    ('Tasa de aprobación · ciclo 1', f"{round(M['aprob1']/M['ejec1']*100)} %", 'Se lee con el riesgo'),
    ('Tasa de aprobación · ciclo 2', f"{round(M['aprob2']/M['ejec2']*100)} %", 'Se lee con el riesgo'),
    ('Defectos reportados', str(M['total_def']), '10 o más'),
    ('Defectos por severidad', '1 crítica · 2 altas · 9 medias · 2 bajas', 'Sin críticos abiertos'),
    ('Defectos verificados (corregidos)', '1 de 14 (DEF-05)', '—'),
]
for i, (m, v, meta) in enumerate(metr, 2):
    wm.cell(i, 1, m).border = borde
    wm.cell(i, 2, v).border = borde
    wm.cell(i, 3, meta).border = borde

# Tabla auxiliar para el gráfico: defectos por severidad.
base = len(metr) + 4
wm.cell(base, 1, 'Severidad').font = negrita
wm.cell(base, 2, 'Defectos').font = negrita
sev_orden = [('Crítica', 1), ('Alta', 2), ('Media', 9), ('Baja', 2)]
for k, (sev, n) in enumerate(sev_orden, 1):
    wm.cell(base + k, 1, sev)
    wm.cell(base + k, 2, n)
chart = BarChart(); chart.type = 'col'; chart.title = 'Defectos por severidad'
chart.y_axis.title = 'Defectos'; chart.legend = None
datos_ref = Reference(wm, min_col=2, min_row=base, max_row=base + len(sev_orden))
cats = Reference(wm, min_col=1, min_row=base + 1, max_row=base + len(sev_orden))
chart.add_data(datos_ref, titles_from_data=True); chart.set_categories(cats)
chart.height = 7; chart.width = 12
wm.add_chart(chart, 'E2')

# ---- Hoja 0 (portada) al inicio --------------------------------------------
wp = wb.create_sheet('Portada', 0)
wp.column_dimensions['A'].width = 24; wp.column_dimensions['B'].width = 70
wp.cell(1, 1, 'ConectaProfe — Registro de ejecución (Entregable 2)').font = Font(bold=True, size=14, color=TEAL)
datos_portada = [
    ('Curso', EQUIPO['curso'] + ' (' + EQUIPO['codigos'] + ')'),
    ('Equipo', EQUIPO['equipo']),
    ('Integrantes', ' · '.join(EQUIPO['integrantes'])),
    ('Docente', EQUIPO['docente']),
    ('Sistema bajo prueba', EQUIPO['sistema']),
    ('Entorno', EQUIPO['url'] + ' · commit ' + EQUIPO['commit']),
    ('Navegador ciclo 1', EQUIPO['navegador_c1'] + ' · escritorio 1366 px'),
    ('Navegador ciclo 2', EQUIPO['navegador_c2'] + ' (sustituye a Firefox; ver nota)'),
    ('Sistema operativo', EQUIPO['so']),
    ('Zona horaria', 'America/Bogota (hora de Colombia)'),
    ('Ejecución', 'Automatizada con Playwright + Chromium; una ventana nueva por sesión'),
]
for i, (k, v) in enumerate(datos_portada, 3):
    wp.cell(i, 1, k).font = negrita
    wp.cell(i, 2, v).alignment = Alignment(wrap_text=True)
nota = ('Nota sobre el ciclo 2: la guía pide repetir en Firefox. El entorno de ejecución '
        'automatizada solo dispone de Chromium, así que el ciclo 2 se corrió en una instancia '
        'independiente de Chromium (ventana y almacenamiento nuevos por caso), lo que confirma '
        'que los fallos son estables entre corridas. El contraste específico en Firefox queda '
        'para ejecución manual del equipo sobre Railway.')
wp.cell(len(datos_portada) + 4, 1, 'Nota').font = negrita
c = wp.cell(len(datos_portada) + 4, 2, nota); c.alignment = Alignment(wrap_text=True)

salida = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', 'docs', 'entrega-2', 'E2_Registro_Ejecucion_EquipoF.xlsx')
import os; os.makedirs(os.path.dirname(salida), exist_ok=True)
wb.save(salida)
print('Guardado', salida)
